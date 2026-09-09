import pandas as pd
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import streamlit as st
# Load the data
df1 = pd.read_excel('q1_2026_school_charges_cleaned.xlsx')
df2 = pd.read_excel('q2_2026_school_charges_cleaned.xlsx')
df = pd.concat([df1, df2], ignore_index=True)

print("Creating comprehensive Q1+Q2 2026 report...")

# Create summary tables
# 1. Overall summary
overall_summary = pd.DataFrame({
    'Category': ['Total Students', 'Indigene', 'Non-Indigene', 'Science', 'Non-Science'],
    'Count': [
        df['NO OF STUDENTS'].sum(),
        df[df['STATUS'] == 'INDIGENE']['NO OF STUDENTS'].sum(),
        df[df['STATUS'] == 'NON-INDIGENE']['NO OF STUDENTS'].sum(),
        df[df['SCIENCE/NON-SCIENCE'] == 'SCIENCE']['NO OF STUDENTS'].sum(),
        df[df['SCIENCE/NON-SCIENCE'] == 'NON-SCIENCE']['NO OF STUDENTS'].sum()
    ]
})

# 2. By Level
level_summary = df.groupby('LEVEL')['NO OF STUDENTS'].sum().reset_index()
level_summary.columns = ['Level', 'Total Students']

# 3. By Status
status_summary = df.groupby('STATUS')['NO OF STUDENTS'].sum().reset_index()
status_summary.columns = ['Status', 'Total Students']

# 4. By Science/Non-Science
science_summary = df.groupby('SCIENCE/NON-SCIENCE')['NO OF STUDENTS'].sum().reset_index()
science_summary.columns = ['Program Type', 'Total Students']

# 5. By Level and Status
level_status_summary = df.groupby(['LEVEL', 'STATUS'])['NO OF STUDENTS'].sum().unstack(fill_value=0).reset_index()
level_status_summary.columns.name = None
level_status_summary.columns = ['Level', 'Indigene', 'Non-Indigene']
level_status_summary['Total'] = level_status_summary['Indigene'] + level_status_summary['Non-Indigene']

# 6. By Level and Science/Non-Science
level_science_summary = df.groupby(['LEVEL', 'SCIENCE/NON-SCIENCE'])['NO OF STUDENTS'].sum().unstack(fill_value=0).reset_index()
level_science_summary.columns.name = None
level_science_summary.columns = ['Level', 'Non-Science', 'Science']
level_science_summary['Total'] = level_science_summary['Non-Science'] + level_science_summary['Science']

# 7. By Status and Science/Non-Science
status_science_summary = df.groupby(['STATUS', 'SCIENCE/NON-SCIENCE'])['NO OF STUDENTS'].sum().unstack(fill_value=0).reset_index()
status_science_summary.columns.name = None
status_science_summary.columns = ['Status', 'Non-Science', 'Science']
status_science_summary['Total'] = status_science_summary['Non-Science'] + status_science_summary['Science']

# 8. Detailed breakdown by Level, Status, and Science/Non-Science
detailed_summary = df.pivot_table(
    values='NO OF STUDENTS',
    index='LEVEL',
    columns=['STATUS', 'SCIENCE/NON-SCIENCE'],
    aggfunc='sum',
    fill_value=0
).reset_index()
detailed_summary.columns.name = None
detailed_summary.columns = ['Level', 'Indigene_Non-Science', 'Indigene_Science', 
                            'Non-Indigene_Non-Science', 'Non-Indigene_Science']
detailed_summary['Total'] = detailed_summary[['Indigene_Non-Science', 'Indigene_Science', 
                                              'Non-Indigene_Non-Science', 'Non-Indigene_Science']].sum(axis=1)

# 9. Monthly summary
monthly_summary = df.groupby('MONTH')['NO OF STUDENTS'].sum().reset_index()
monthly_summary.columns = ['Month', 'Total Students']

# 10. Monthly by Level
monthly_level_summary = df.groupby(['MONTH', 'LEVEL'])['NO OF STUDENTS'].sum().unstack(fill_value=0).reset_index()
monthly_level_summary.columns.name = None

# Save to Excel with formatting
excel_path = r"C:\Users\ALH MALIK TOPE\PycharmProjects\PythonProject\q1_q2_2026_comprehensive_report.xlsx"
with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
    overall_summary.to_excel(writer, sheet_name='Overall Summary', index=False)
    level_summary.to_excel(writer, sheet_name='By Level', index=False)
    status_summary.to_excel(writer, sheet_name='By Status', index=False)
    science_summary.to_excel(writer, sheet_name='By Program Type', index=False)
    level_status_summary.to_excel(writer, sheet_name='Level & Status', index=False)
    level_science_summary.to_excel(writer, sheet_name='Level & Program', index=False)
    status_science_summary.to_excel(writer, sheet_name='Status & Program', index=False)
    detailed_summary.to_excel(writer, sheet_name='Detailed Breakdown', index=False)
    monthly_summary.to_excel(writer, sheet_name='Monthly Summary', index=False)
    monthly_level_summary.to_excel(writer, sheet_name='Monthly by Level', index=False)
    df.to_excel(writer, sheet_name='Raw Data', index=False)

print(f"Excel report saved to: {excel_path}")

# Format Excel file
wb = load_workbook(excel_path)

# Define styles
header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
center_alignment = Alignment(horizontal='center', vertical='center')
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Format each sheet
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    
    # Format header row
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        cell.border = border
    
    # Format data cells
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = center_alignment
            cell.border = border
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 30)
        ws.column_dimensions[column_letter].width = adjusted_width

wb.save(excel_path)
print("Excel formatting completed")

# Create PDF report with visualizations
pdf_path = r"C:\Users\ALH MALIK TOPE\PycharmProjects\PythonProject\q1_q2_2026_comprehensive_report.pdf"

with PdfPages(pdf_path) as pdf:
    # Page 1: Overall Summary
    fig, ax = plt.subplots(figsize=(11, 8))
    ax.axis('tight')
    ax.axis('off')
    
    # Create table
    table_data = []
    table_data.append(['Q1+Q2 2026 SCHOOL FEES ANALYSIS REPORT', ''])
    table_data.append(['', ''])
    table_data.append(['OVERALL SUMMARY', ''])
    table_data.append(['', ''])
    table_data.append(['Category', 'Count'])
    
    for _, row in overall_summary.iterrows():
        table_data.append([row['Category'], f"{row['Count']:,}"])
    
    table = ax.table(cellText=table_data, cellLoc='center', loc='center')
    table.auto_set_font_size(False)
    table.set_fontsize(12)
    table.scale(1, 2)
    
    # Style the header
    for i in range(5):
        for j in range(2):
            table[(i, j)].set_facecolor('#4472C4')
            table[(i, j)].set_text_props(weight='bold', color='white')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 2: By Level
    fig, ax = plt.subplots(figsize=(11, 8))
    ax.axis('tight')
    ax.axis('off')
    
    table_data = []
    table_data.append(['STUDENTS BY LEVEL', ''])
    table_data.append(['', ''])
    table_data.append(['Level', 'Total Students'])
    
    for _, row in level_summary.iterrows():
        table_data.append([row['Level'], f"{row['Total Students']:,}"])
    
    table = ax.table(cellText=table_data, cellLoc='center', loc='center')
    table.auto_set_font_size(False)
    table.set_fontsize(12)
    table.scale(1, 1.5)
    
    for i in range(3):
        for j in range(2):
            table[(i, j)].set_facecolor('#4472C4')
            table[(i, j)].set_text_props(weight='bold', color='white')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 3: By Status
    fig, ax = plt.subplots(figsize=(11, 8))
    ax.axis('tight')
    ax.axis('off')
    
    table_data = []
    table_data.append(['STUDENTS BY STATUS (INDIGENE VS NON-INDIGENE)', ''])
    table_data.append(['', ''])
    table_data.append(['Status', 'Total Students'])
    
    for _, row in status_summary.iterrows():
        table_data.append([row['Status'], f"{row['Total Students']:,}"])
    
    table = ax.table(cellText=table_data, cellLoc='center', loc='center')
    table.auto_set_font_size(False)
    table.set_fontsize(12)
    table.scale(1, 1.5)
    
    for i in range(3):
        for j in range(2):
            table[(i, j)].set_facecolor('#4472C4')
            table[(i, j)].set_text_props(weight='bold', color='white')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 4: By Program Type
    fig, ax = plt.subplots(figsize=(11, 8))
    ax.axis('tight')
    ax.axis('off')
    
    table_data = []
    table_data.append(['STUDENTS BY PROGRAM TYPE (SCIENCE VS NON-SCIENCE)', ''])
    table_data.append(['', ''])
    table_data.append(['Program Type', 'Total Students'])
    
    for _, row in science_summary.iterrows():
        table_data.append([row['Program Type'], f"{row['Total Students']:,}"])
    
    table = ax.table(cellText=table_data, cellLoc='center', loc='center')
    table.auto_set_font_size(False)
    table.set_fontsize(12)
    table.scale(1, 1.5)
    
    for i in range(3):
        for j in range(2):
            table[(i, j)].set_facecolor('#4472C4')
            table[(i, j)].set_text_props(weight='bold', color='white')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 5: Level & Status
    fig, ax = plt.subplots(figsize=(11, 8))
    ax.axis('tight')
    ax.axis('off')
    
    table_data = []
    table_data.append(['STUDENTS BY LEVEL AND STATUS', '', '', ''])
    table_data.append(['', '', '', ''])
    table_data.append(['Level', 'Indigene', 'Non-Indigene', 'Total'])
    
    for _, row in level_status_summary.iterrows():
        table_data.append([row['Level'], f"{row['Indigene']:,}", 
                          f"{row['Non-Indigene']:,}", f"{row['Total']:,}"])
    
    table = ax.table(cellText=table_data, cellLoc='center', loc='center')
    table.auto_set_font_size(False)
    table.set_fontsize(11)
    table.scale(1, 1.5)
    
    for i in range(3):
        for j in range(4):
            table[(i, j)].set_facecolor('#4472C4')
            table[(i, j)].set_text_props(weight='bold', color='white')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 6: Level & Program
    fig, ax = plt.subplots(figsize=(11, 8))
    ax.axis('tight')
    ax.axis('off')
    
    table_data = []
    table_data.append(['STUDENTS BY LEVEL AND PROGRAM TYPE', '', '', ''])
    table_data.append(['', '', '', ''])
    table_data.append(['Level', 'Non-Science', 'Science', 'Total'])
    
    for _, row in level_science_summary.iterrows():
        table_data.append([row['Level'], f"{row['Non-Science']:,}", 
                          f"{row['Science']:,}", f"{row['Total']:,}"])
    
    table = ax.table(cellText=table_data, cellLoc='center', loc='center')
    table.auto_set_font_size(False)
    table.set_fontsize(11)
    table.scale(1, 1.5)
    
    for i in range(3):
        for j in range(4):
            table[(i, j)].set_facecolor('#4472C4')
            table[(i, j)].set_text_props(weight='bold', color='white')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 7: Status & Program
    fig, ax = plt.subplots(figsize=(11, 8))
    ax.axis('tight')
    ax.axis('off')
    
    table_data = []
    table_data.append(['STUDENTS BY STATUS AND PROGRAM TYPE', '', '', ''])
    table_data.append(['', '', '', ''])
    table_data.append(['Status', 'Non-Science', 'Science', 'Total'])
    
    for _, row in status_science_summary.iterrows():
        table_data.append([row['Status'], f"{row['Non-Science']:,}", 
                          f"{row['Science']:,}", f"{row['Total']:,}"])
    
    table = ax.table(cellText=table_data, cellLoc='center', loc='center')
    table.auto_set_font_size(False)
    table.set_fontsize(12)
    table.scale(1, 1.5)
    
    for i in range(3):
        for j in range(4):
            table[(i, j)].set_facecolor('#4472C4')
            table[(i, j)].set_text_props(weight='bold', color='white')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 8: Detailed Breakdown
    fig, ax = plt.subplots(figsize=(11, 8))
    ax.axis('tight')
    ax.axis('off')
    
    table_data = []
    table_data.append(['DETAILED BREAKDOWN BY LEVEL, STATUS, AND PROGRAM TYPE', '', '', '', '', ''])
    table_data.append(['', '', '', '', '', ''])
    table_data.append(['Level', 'Indigene\nNon-Science', 'Indigene\nScience', 
                      'Non-Indigene\nNon-Science', 'Non-Indigene\nScience', 'Total'])
    
    for _, row in detailed_summary.iterrows():
        table_data.append([row['Level'], f"{row['Indigene_Non-Science']:,}", 
                          f"{row['Indigene_Science']:,}", 
                          f"{row['Non-Indigene_Non-Science']:,}", 
                          f"{row['Non-Indigene_Science']:,}", 
                          f"{row['Total']:,}"])
    
    table = ax.table(cellText=table_data, cellLoc='center', loc='center')
    table.auto_set_font_size(False)
    table.set_fontsize(10)
    table.scale(1, 1.5)
    
    for i in range(3):
        for j in range(6):
            table[(i, j)].set_facecolor('#4472C4')
            table[(i, j)].set_text_props(weight='bold', color='white')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 9: Monthly Summary
    fig, ax = plt.subplots(figsize=(11, 8))
    ax.axis('tight')
    ax.axis('off')
    
    table_data = []
    table_data.append(['MONTHLY SUMMARY', ''])
    table_data.append(['', ''])
    table_data.append(['Month', 'Total Students'])
    
    for _, row in monthly_summary.iterrows():
        table_data.append([row['Month'], f"{row['Total Students']:,}"])
    
    table = ax.table(cellText=table_data, cellLoc='center', loc='center')
    table.auto_set_font_size(False)
    table.set_fontsize(12)
    table.scale(1, 1.5)
    
    for i in range(3):
        for j in range(2):
            table[(i, j)].set_facecolor('#4472C4')
            table[(i, j)].set_text_props(weight='bold', color='white')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 10: Charts - Status Pie Chart
    fig, ax = plt.subplots(figsize=(11, 8))
    status_data = df.groupby('STATUS')['NO OF STUDENTS'].sum()
    colors = ['#FF6B6B', '#4ECDC4']
    wedges, texts, autotexts = ax.pie(status_data, labels=status_data.index, autopct='%1.1f%%',
                                      colors=colors, startangle=90, 
                                      textprops={'fontsize': 14, 'fontweight': 'bold'})
    ax.set_title('Students by Status - Q1+Q2 2026', fontsize=16, fontweight='bold', pad=20)
    
    for autotext in autotexts:
        autotext.set_color('white')
        autotext.set_fontsize(12)
        autotext.set_fontweight('bold')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 11: Charts - Level Pie Chart
    fig, ax = plt.subplots(figsize=(11, 8))
    level_data = df.groupby('LEVEL')['NO OF STUDENTS'].sum().sort_values(ascending=False)
    colors = plt.cm.Set3(range(len(level_data)))
    wedges, texts, autotexts = ax.pie(level_data, labels=level_data.index, autopct='%1.1f%%',
                                      colors=colors, startangle=90, 
                                      textprops={'fontsize': 12, 'fontweight': 'bold'})
    ax.set_title('Students by Level - Q1+Q2 2026', fontsize=16, fontweight='bold', pad=20)
    
    for autotext in autotexts:
        autotext.set_color('white')
        autotext.set_fontsize(10)
        autotext.set_fontweight('bold')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()
    
    # Page 12: Charts - Science Pie Chart
    fig, ax = plt.subplots(figsize=(11, 8))
    science_data = df.groupby('SCIENCE/NON-SCIENCE')['NO OF STUDENTS'].sum()
    colors = ['#45B7D1', '#96CEB4']
    wedges, texts, autotexts = ax.pie(science_data, labels=science_data.index, autopct='%1.1f%%',
                                      colors=colors, startangle=90, 
                                      textprops={'fontsize': 14, 'fontweight': 'bold'})
    ax.set_title('Students by Program Type - Q1+Q2 2026', fontsize=16, fontweight='bold', pad=20)
    
    for autotext in autotexts:
        autotext.set_color('white')
        autotext.set_fontsize(12)
        autotext.set_fontweight('bold')
    
    pdf.savefig(fig, bbox_inches='tight')
    plt.close()

print(f"PDF report saved to: {pdf_path}")
print("Report generation completed successfully!")
